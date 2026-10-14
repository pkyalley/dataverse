#!/usr/bin/env python3
"""Independent workbook-based reproduction of CMS NHE24 figures.

Uses archived repository sources and creates corrected/reframed outputs.
Does not modify published figures or claim causal interpretation.
"""
from pathlib import Path
import json, math, csv
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scipy.stats import pearsonr
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
DATA=ROOT.parent/"data"
OUT=ROOT/"results"
FIG=ROOT/"figures"
OUT.mkdir(exist_ok=True)
FIG.mkdir(exist_ok=True)
FILES={
 "t02":"Table 02 National Health Expenditures, Aggregate and Per Capita Amounts, by Type of Expenditure.xlsx",
 "t03":"Table 03 National Health Expenditures, by Source of Funds.xlsx",
 "t07":"Table 07 Hospital Care Expenditures.xlsx"
}
def workbook_rows(name):
    wb=load_workbook(DATA/name,data_only=True,read_only=True)
    ws=wb.worksheets[0]
    rows=list(ws.iter_rows(values_only=True))
    wb.close()
    return rows

def table_by_year(name,header_row=2,start_row=4,stop_label="Average Annual Percent Change"):
    rows=workbook_rows(name)
    headers=rows[header_row-1]
    cols={j:int(v) for j,v in enumerate(headers) if isinstance(v,(int,float)) and float(v).is_integer() and 1960<=int(v)<=2024}
    assert len(cols)>=8,(name,cols)
    output={}
    for row in rows[start_row-1:]:
        label=str(row[0] or "").strip()
        if stop_label.lower() in " ".join(str(v or "") for v in row[:2]).lower():break
        if not label:continue
        vals={}
        for j,y in cols.items():
            v=row[j] if j<len(row) else None
            if isinstance(v,(int,float)) and np.isfinite(v): vals[y]=float(v)
        if vals:
            assert label not in output,("duplicate label",label)
            output[label]=pd.Series(vals,dtype=float).sort_index()
    return output

def select(data,phrase):
    keys=[k for k in data if phrase.casefold() in k.casefold()]
    assert len(keys)==1,(phrase,keys)
    return data[keys[0]]

def chart(fig,name):
    fig.tight_layout()
    fig.savefig(FIG/name,dpi=175,bbox_inches="tight")
    plt.close(fig)

def main():
    t2=table_by_year(FILES["t02"])
    t3=table_by_year(FILES["t03"])
    hospital=select(t2,"Hospital Care")
    physician=select(t2,"Physician and Clinical")
    rx=select(t2,"Prescription Drugs")
    nursing=select(t2,"Nursing Care Facilities")
    home=select(t2,"Home Health Care")
    nhe=select(t2,"National Health Expenditures")
    oop=select(t3,"Out of pocket")
    private=select(t3,"Private Health Insurance")
    medicare=select(t3,"Medicare")
    medicaid=select(t3,"Medicaid")
    nhe3=select(t3,"National Health Expenditures")
    assert abs(nhe[2024]-nhe3[2024])<0.1,"Table 02/03 NHE mismatch"
    # The archived Summary CSV is independently parsed by reproduce.py.
    # Table 02 government administration and non-medical insurance are distinct subcategories.
    gov=select(t2,"Government Administration")
    nonmed=select(t2,"Non-Medical Insurance Expenditures")
    admin=gov.add(nonmed,fill_value=0)
    years=sorted(set(nhe.index)&set(admin.index))
    assert abs(admin[2024]-372.1)<0.3,(admin[2024],"Expected 372.1B within rounding")
    assert abs(hospital[2024]-1634.7)<0.2,(hospital[2024],"Hospital 2024 unexpected")
    stats={}
    for name,s in [("hospital",hospital),("physician_clinical",physician),("rx_drugs",rx),("nursing",nursing),("home_health",home)]:
        stats[name]={"1990_billion":float(s[1990]),"2024_billion":float(s[2024]),
                     "nominal_growth_pct":float(100*(s[2024]/s[1990]-1))}
    payers={"Medicare":medicare,"Private insurance":private,"Medicaid":medicaid,"Out of pocket":oop}
    selected_sum=sum(s[2024] for s in payers.values())
    stats["payer_2024"]={k:{"billions":float(s[2024]),
                            "share_four_selected_pct":float(s[2024]/selected_sum*100),
                            "share_total_nhe_pct":float(s[2024]/nhe[2024]*100)}
                         for k,s in payers.items()}
    stats["payer_2024"]["Other NHE (remainder)"]={"billions":float(nhe[2024]-selected_sum),
                      "share_total_nhe_pct":float(100*(nhe[2024]-selected_sum)/nhe[2024])}
    stats["correlation_1980_2024"]={}
    # Table 02 has 1980, 1990, 2000 and annual 2001-2024, NOT every year 1980-1999.
    # Only common available years may be used; do not fabricate intermediate years.
    y=[v for v in years if v>=1980]
    stats["correlation_1980_2024"]["n_observed_years"]=len(y)
    stats["correlation_1980_2024"]["pearson_r_available_years"]=float(pearsonr([nhe[v] for v in y],[admin[v] for v in y]).statistic)
    stats["source_years"]=years
    (OUT/"workbook_reproduction.json").write_text(json.dumps(stats,indent=2)+"\n")
    with (OUT/"category_comparison.csv").open("w",newline="") as f:
        w=csv.writer(f);w.writerow(["category","1990_billion","2024_billion","nominal_growth_pct"])
        for name in ("hospital","physician_clinical","rx_drugs","nursing","home_health"):
            x=stats[name];w.writerow([name,x["1990_billion"],x["2024_billion"],x["nominal_growth_pct"]])
    with (OUT/"payer_2024.csv").open("w",newline="") as f:
        w=csv.writer(f);w.writerow(["category","billions","share_total_nhe_pct","share_four_selected_pct"])
        for k,x in stats["payer_2024"].items():w.writerow([k,x["billions"],x["share_total_nhe_pct"],x.get("share_four_selected_pct","")])
    # Corrected Figure 3: share of TOTAL NHE, including all omitted categories.
    common=sorted(set.intersection(*(set(s.index) for s in [nhe,*payers.values()])))
    fig,ax=plt.subplots(figsize=(11,6))
    labels=list(payers)+["Other NHE"]
    values=[[float(payers[k][yr]/nhe[yr]*100) for yr in common] for k in payers]
    values.append([float(100-sum(v[i] for v in values)) for i in range(len(common))])
    ax.stackplot(common,*values,labels=labels)
    ax.set(xlabel="Year (only years present in archived table)",ylabel="Share of total NHE (%)",
           title="Corrected payer distribution, including other NHE categories",ylim=(0,100))
    ax.legend(loc="upper left",fontsize=8,ncol=3)
    chart(fig,"corrected_fig3_payer_mix.png")
    # Corrected Figure 4: service category amounts; no claim these cover all NHE.
    svc={"Hospital care":hospital[2024],"Physician and clinical":physician[2024],
         "Prescription drugs":rx[2024],"Government admin + non-medical insurance":admin[2024],
         "Nursing care facilities":nursing[2024],"Home health care":home[2024]}
    fig,ax=plt.subplots(figsize=(11,6))
    keys=sorted(svc,key=svc.get)
    ax.barh(keys,[svc[k] for k in keys])
    ax.set(xlabel="2024 expenditures (billions USD)",title="Selected CMS expenditure categories (not all NHE)")
    chart(fig,"corrected_fig4_services.png")
    # Corrected Figure 5: observed year index, no interpolation.
    fig,ax=plt.subplots(figsize=(11,6))
    for k,s in [("Admin and non-medical insurance",admin),("Hospital",hospital),("Physician and clinical",physician),("Prescription drugs",rx)]:
        ys=[yr for yr in s.index if yr>=1990]
        ax.plot(ys,[s[yr]/s[1990]*100 for yr in ys],label=k)
    ax.set(xlabel="Year (observed archived years)",ylabel="Nominal index (1990=100)",title="Spending growth indices")
    ax.legend()
    chart(fig,"corrected_fig5_growth_index.png")
    # Corrected Figure 7: available years only; correlation of nominal levels.
    series={"NHE":nhe,"Admin":admin,"Hospital":hospital,"Physician":physician,"Rx drugs":rx}
    common=sorted(set.intersection(*(set(s.index) for s in series.values())))
    df=pd.DataFrame({k:[s[y] for y in common] for k,s in series.items()},index=common)
    fig,ax=plt.subplots(figsize=(8,7))
    corr=df.corr()
    im=ax.imshow(corr.values,vmin=-1,vmax=1,cmap="coolwarm")
    ax.set_xticks(range(len(corr)),labels=corr.columns,rotation=35,ha="right")
    ax.set_yticks(range(len(corr)),labels=corr.index)
    for i in range(len(corr)):
        for j in range(len(corr)):ax.text(j,i,f"{corr.iloc[i,j]:.3f}",ha="center",va="center",fontsize=10)
    fig.colorbar(im,ax=ax)
    ax.set_title("Nominal spending-level correlations (descriptive only)")
    chart(fig,"corrected_fig7_correlation.png")
    # Corrected Figure 8: explicitly observed 1980+ years.
    fig,ax=plt.subplots(figsize=(9,6))
    ax.scatter([nhe[v] for v in y],[admin[v] for v in y])
    r=stats["correlation_1980_2024"]["pearson_r_available_years"]
    ax.set(xlabel="NHE (billions USD)",ylabel="Gov admin + non-medical insurance (billions USD)",
           title=f"Observed-year nominal correlation, r={r:.3f}; n={len(y)} (not causal)")
    chart(fig,"corrected_fig8_scatter.png")
    # Corrected Figure 9: complete NHE denominator.
    fig,ax=plt.subplots(figsize=(9,7))
    vals=[stats["payer_2024"][k]["billions"] for k in list(payers)+["Other NHE (remainder)"]]
    ax.pie(vals,labels=labels,autopct="%1.1f%%",startangle=90)
    ax.set_title("2024 shares of total NHE (including all remaining categories)")
    chart(fig,"corrected_fig9_payer_donut.png")
    # Corrected Figure 10: scenario, no statistical confidence band.
    fy=list(range(2024,2034))
    fig,ax=plt.subplots(figsize=(10,6))
    ax.plot([v for v in nhe.index if v>=2000],[nhe[v]/1000 for v in nhe.index if v>=2000],label="Archived observed NHE")
    ax.plot(fy,[nhe[2024]*1.058**(v-2024)/1000 for v in fy],ls="--",label="Illustrative constant 5.8% scenario")
    ax.set(xlabel="Year",ylabel="NHE (trillions USD)",title="Illustrative NHE trajectory; NOT official CMS forecast")
    ax.legend()
    chart(fig,"corrected_fig10_scenario.png")
    print(json.dumps(stats,indent=2))
    print("Wrote 7 figures and 3 workbook-derived data files.")

if __name__=="__main__":main()
