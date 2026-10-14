#!/usr/bin/env python3
"""Inspect CMS NHE24 Table 02/03/07 workbook layouts without modifying originals.

Requires openpyxl. Produces a human-reviewable inventory, not corrected values.
Usage: python nhe-analysis/audit-2026/inspect_workbooks.py
"""
from pathlib import Path
import csv, hashlib, json
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
DATA=HERE.parent/"data"
OUT=HERE/"results"
SOURCES={
 "02":"Table 02 National Health Expenditures, Aggregate and Per Capita Amounts, by Type of Expenditure.xlsx",
 "03":"Table 03 National Health Expenditures, by Source of Funds.xlsx",
 "07":"Table 07 Hospital Care Expenditures.xlsx"
}

def inspect(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    report={"file":path.name,"sha256":hashlib.sha256(path.read_bytes()).hexdigest(),"sheets":[]}
    for ws in wb.worksheets:
        rows=ws.iter_rows(values_only=True)
        candidates=[]
        for idx,row in enumerate(rows,1):
            years=[(j,int(v)) for j,v in enumerate(row) if isinstance(v,(int,float)) and 1960<=v<=2030 and float(v).is_integer()]
            if len(years)>=4:
                candidates.append({"row":idx,"year_columns_1_based":[j+1 for j,_ in years],"years":[v for _,v in years]})
            if idx>=25:break
        report["sheets"].append({"sheet":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,"candidate_year_headers":candidates})
    wb.close()
    return report

def main():
    OUT.mkdir(exist_ok=True)
    reports=[inspect(DATA/name) for name in SOURCES.values()]
    (OUT/"workbook_inventory.json").write_text(json.dumps(reports,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(reports,indent=2))
    print("Next: manually verify the header row, dollar units, row labels, and year coverage before extracting expenditure series.")

if __name__=="__main__":main()
