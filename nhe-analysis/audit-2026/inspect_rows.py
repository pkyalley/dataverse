#!/usr/bin/env python3
from pathlib import Path
from openpyxl import load_workbook
import json
ROOT=Path(__file__).resolve().parent
D=ROOT.parent/"data"
names={
"02":"Table 02 National Health Expenditures, Aggregate and Per Capita Amounts, by Type of Expenditure.xlsx",
"03":"Table 03 National Health Expenditures, by Source of Funds.xlsx",
"07":"Table 07 Hospital Care Expenditures.xlsx"
}
out={}
for key,name in names.items():
 w=load_workbook(D/name,read_only=True,data_only=True)
 out[key]={}
 for s in w.worksheets:
  entries=[]
  for i,row in enumerate(s.iter_rows(values_only=True),1):
   if i>75:break
   vals=[{"col":j+1,"v":str(v)[:90]} for j,v in enumerate(row) if v is not None]
   if vals:entries.append({"row":i,"cells":vals[:12]})
  out[key][s.title]=entries
 w.close()
p=ROOT/"results"/"workbook_first75rows.json"
p.parent.mkdir(exist_ok=True)
p.write_text(json.dumps(out,indent=2))
for k,v in out.items():
 print("\nTABLE",k)
 for sheet,rows in v.items():
  print("SHEET",sheet)
  for row in rows[:45]:print(row["row"]," | ".join(str(x["col"])+":"+x["v"][:38] for x in row["cells"][:6]))
